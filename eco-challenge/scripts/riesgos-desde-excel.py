#!/usr/bin/env python3
"""
Genera el SQL del tema "Riesgos Críticos" a partir del Excel de controles
(hojas "Controles por RC" y "Preguntas de verificación").

    python3 -I scripts/riesgos-desde-excel.py Controles.xlsx > riesgos.sql

El SQL resultante crea:
  - dos ámbitos (grupos de hasta 7 riesgos, lo que cabe en un celular),
  - un "contenedor" por riesgo crítico (color, icono, código y nombre corto),
  - una tarjeta por control (respuesta = su riesgo crítico), con una
    explicación que incluye el tipo de control y una pregunta de verificación,
  - un banco de preguntas por grupo.

Un control cuyo texto se repite en dos riesgos del MISMO grupo se excluye,
porque tendría más de una respuesta correcta.

IMPORTANTE: el Excel y el SQL generado pueden ser confidenciales; no los
suba a un repositorio público. Aplíquelos directamente en Supabase.
Requiere openpyxl.
"""
import re
import sys
import unicodedata
import uuid
from difflib import SequenceMatcher
from collections import OrderedDict, defaultdict

import openpyxl

NS = uuid.UUID('5e3c1d2a-6f7b-4c8d-9e0f-1a2b3c4d5e6f')

# Presentación de cada riesgo en el juego (nombre corto, icono y color).
RC_META = {
    'RC5':  ('confinados',    'Espacios confinados',        '🕳️', '#5B3A8C'),
    'RC6':  ('agua',          'Cuerpos de agua',            '🌊', '#1E63C6'),
    'RC7':  ('equipo_minero', 'Equipo minero',              '🚜', '#B26A00'),
    'RC11': ('explosivos',    'Explosivos',                 '💥', '#C62828'),
    'RC12': ('atropello',     'Atropello',                  '🚶', '#2E7D32'),
    'RC13': ('caida_rocas',   'Caída de rocas',             '🪨', '#6D4C41'),
    'RC14': ('caliente',      'Trabajos en caliente',       '🔥', '#E65100'),
    'RC15': ('izaje',         'Maniobras de izaje',         '🏗️', '#F2C200'),
    'RC16': ('energia',       'Energía peligrosa',          '⚙️', '#455A64'),
    'RC20': ('atrapamiento',  'Atrapamiento / aplastamiento', '✋', '#AD1457'),
    'RC21': ('altura',        'Trabajos en altura',         '🧗', '#00838F'),
    'RC22': ('electrico',     'Energía eléctrica',          '⚡', '#283593'),
    'RC23': ('rayos',         'Rayos / zonas remotas',      '🌩️', '#37474F'),
    'RC28': ('transporte',    'Transporte de personal',     '🚌', '#2E9E44'),
}
DARK_TEXT = {'#F2C200'}

GROUPS = [
    ('rc_grupo_a', 'Riesgos críticos · Grupo A', 1),
    ('rc_grupo_b', 'Riesgos críticos · Grupo B', 2),
]


def q(v):
    return "'" + str(v).replace("'", "''") + "'"


def rc_num(code):
    return int(re.sub(r'\D', '', code))


def main(path):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ctrl_rows = [r for r in wb['Controles por RC'].iter_rows(min_row=2, values_only=True) if r and r[0]]
    questions = defaultdict(list)
    for r in wb['Preguntas de verificación'].iter_rows(min_row=2, values_only=True):
        if r and r[0] and r[3] and r[4]:
            questions[(r[0].strip(), r[3].strip())].append(re.sub(r'^\s*\d+\.\s*', '', str(r[4]).strip()))

    rcs = OrderedDict()
    for code, rc_name, tipo, _n, control, _nq in ctrl_rows:
        rcs.setdefault(code.strip(), {'name': rc_name.strip(), 'controls': []})
        rcs[code.strip()]['controls'].append((tipo.strip(), control.strip()))

    missing = [c for c in rcs if c not in RC_META]
    if missing:
        sys.exit(f'Faltan metadatos de presentación para: {missing}')

    codes = sorted(rcs, key=rc_num)
    half = (len(codes) + 1) // 2
    group_of = {c: GROUPS[0] if i < half else GROUPS[1] for i, c in enumerate(codes)}
    if half > 7 or len(codes) - half > 7:
        sys.exit('Cada grupo admite como máximo 7 riesgos (contenedores en el celular).')

    out = ['-- Generado por scripts/riesgos-desde-excel.py. CONTENIDO CONFIDENCIAL: no publicar.', 'begin;',
           '-- Se registra como validado por el administrador del sistema.',
           "select set_config('request.jwt.claims', json_build_object('sub', "
           "(select user_id from public.staff where role = 'admin' order by created_at limit 1), "
           "'role', 'authenticated')::text, true);"]
    for scope, name, sort in GROUPS:
        members = [c for c in codes if group_of[c][0] == scope]
        desc = f'{len(members)} riesgos: ' + ', '.join(members)
        out.append(
            'insert into public.scopes (code, theme, name, description, prompt, sort_order) values '
            f"({q(scope)}, 'riesgos', {q(name)}, {q(desc)}, {q('Arrastra cada control a su riesgo crítico.')}, {10 + sort}) "
            'on conflict (code) do update set name = excluded.name, description = excluded.description, '
            'prompt = excluded.prompt, sort_order = excluded.sort_order;')

    for i, code in enumerate(codes):
        bin_code, short, icon, color = RC_META[code]
        scope = group_of[code][0]
        fg = '#1D2630' if color in DARK_TEXT else '#FFFFFF'
        out.append(
            'insert into public.bins (scope, code, color_name, label, color_hex, text_hex, icon, description, sort_order) values '
            f'({q(scope)}, {q(bin_code)}, {q(code)}, {q(short)}, {q(color)}, {q(fg)}, {q(icon)}, {q(rcs[code]["name"])}, {i + 1}) '
            'on conflict (scope, code) do update set color_name = excluded.color_name, label = excluded.label, '
            'color_hex = excluded.color_hex, text_hex = excluded.text_hex, icon = excluded.icon, '
            'description = excluded.description, sort_order = excluded.sort_order;')

    # Detección de controles ambiguos dentro de cada grupo: textos iguales o
    # casi iguales (p. ej. variantes de "Aislamiento, bloqueo...") en riesgos
    # distintos tendrían más de una respuesta correcta.
    def norm(t):
        t = unicodedata.normalize('NFD', t.lower())
        t = ''.join(ch for ch in t if unicodedata.category(ch) != 'Mn')
        return ' '.join(w for w in re.findall(r'[a-z0-9]+', t) if w not in {'y', 'de', 'del', 'la', 'el', 'en', 'para', 'los', 'las'})

    seen = defaultdict(set)
    entries = defaultdict(list)
    for code in codes:
        for _tipo, control in rcs[code]['controls']:
            entries[group_of[code][0]].append((code, control, norm(control)))
    for scope, lst in entries.items():
        for code, control, n1 in lst:
            seen[(scope, control.lower())].add(code)
            for code2, _c2, n2 in lst:
                if code2 != code and SequenceMatcher(None, n1, n2).ratio() >= 0.8:
                    seen[(scope, control.lower())].add(code2)

    excluded, counts = [], defaultdict(int)
    bank_ids = {scope: uuid.uuid5(NS, 'bank:' + scope) for scope, _n, _s in GROUPS}
    for scope, name, _s in GROUPS:
        out.append(
            'insert into public.question_banks (id, name, scope, description, owner_id) values '
            f"({q(bank_ids[scope])}, {q('Controles · ' + name.split('· ')[-1])}, {q(scope)}, "
            f"{q('Relacionar cada control con su riesgo crítico.')}, null) on conflict (id) do nothing;")

    for scope, _name, _s in GROUPS:
        values, links = [], []
        for code in [c for c in codes if group_of[c][0] == scope]:
            bin_code = RC_META[code][0]
            for tipo, control in rcs[code]['controls']:
                if len(seen[(scope, control.lower())]) > 1:
                    excluded.append(f'{code}: {control}')
                    continue
                item_id = uuid.uuid5(NS, f'item:{code}:{control}')
                qs = questions.get((code, control), [])
                expl = f'Control {tipo.lower()} del {code} – {rcs[code]["name"]}.'
                if qs:
                    expl += f' En campo se verifica, por ejemplo: {qs[0]}'
                expl = expl[:597] + '…' if len(expl) > 600 else expl
                name = control if len(control) <= 160 else control[:157] + '…'
                values.append(f'({q(item_id)},{q(name)},{q(bin_code)},{q(expl)})')
                links.append(q(item_id))
                counts[scope] += 1
        out.append(f'-- {scope}')
        out.append(
            'insert into public.waste_items (id, name, scope, bin_code, image_url, image_source, image_credit, explanation, status, review_note) '
            f"select v.id::uuid, v.name, {q(scope)}, v.bin, '/img/riesgos/control.svg', 'propia', 'Ilustración propia (CC0)', v.expl, 'aprobado', "
            "'Fuente: matriz de riesgos y controles críticos de la empresa' from (values\n" + ',\n'.join(values) +
            '\n) as v(id, name, bin, expl) on conflict (id) do update set name = excluded.name, scope = excluded.scope, '
            'bin_code = excluded.bin_code, explanation = excluded.explanation;')
        out.append(
            f'insert into public.bank_items (bank_id, item_id) select {q(bank_ids[scope])}, x::uuid from unnest(array[' +
            ','.join(links) + ']) as x on conflict do nothing;')
    out.append('commit;')
    print('\n'.join(out))
    print(f'-- Controles por grupo: {dict(counts)}', file=sys.stderr)
    print(f'-- Excluidos por ambigüedad ({len(excluded)}): {excluded}', file=sys.stderr)


if __name__ == '__main__':
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    main(sys.argv[1])
